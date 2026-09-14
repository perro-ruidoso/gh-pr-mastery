"""Build the mastery-objectives workbook (and its markdown mirror) from one source list.

Run: python objectives/build_objectives.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent

UNITS = {
    1: "The machine you work in",
    2: "Issues as units of agent work",
    3: "Branches and the shape of a diff",
    4: "Owning the diff",
    5: "Gates and automation",
    6: "The maintainer seat, PM integration, and capstone",
}

# (id, week, category, objective, bloom, assessed_by, primary_source)
OBJECTIVES = [
    # --- Week 1 ---
    ("M01", 1, "Setup",
     "Join the course GitHub org, authenticate the gh CLI with the scopes the course needs, and instantiate a personal Instance from the flashcards-seed template.",
     "Apply", "A1 (checker)", "GitHub Docs: Creating a repository from a template"),
    ("M02", 1, "The PR model",
     "Explain a pull request as a request to merge one ref into another, naming base and head, and identify both for any open PR.",
     "Understand", "Quiz 1; Self-Check", "GitHub Docs: About pull requests"),
    ("M03", 1, "Diffs",
     "Distinguish two-dot from three-dot diffs and explain why a PR's Files-changed view is computed from the merge base.",
     "Understand", "Quiz 1; Self-Check", "GitHub Docs: About comparing branches in pull requests"),
    ("M04", 1, "Tooling",
     "Inspect issues and PRs from the terminal with gh issue/pr list, view, and diff, including --json field selection.",
     "Apply", "A1 (checker)", "GitHub CLI manual: gh pr, gh issue"),
    ("M05", 1, "Agent config",
     "Configure an Instance for the engineering skills with /setup-matt-pocock-skills: issue tracker, triage label vocabulary, and domain docs.",
     "Apply", "A1 (checker)", "mattpocock/skills: setup-matt-pocock-skills"),
    ("M06", 1, "Evidence",
     "Read the course repo's own issue graph and PR history as an exhibit and describe the workflow that produced it.",
     "Analyze", "A1 (written)", "Course Repo: perro-ruidoso/gh-pr-mastery"),

    # --- Week 2 ---
    ("M07", 2, "Issues",
     "Write an issue body a fresh agent session can execute without asking a question: context, acceptance criteria, and explicit out-of-scope.",
     "Create", "A2 (rubric)", "GitHub Docs: About issues"),
    ("M08", 2, "Issues",
     "Apply labels, issue types, and issue templates so a backlog is filterable and ready work is findable.",
     "Apply", "A2 (checker)", "GitHub Docs: Labels; Issue types; Issue templates"),
    ("M09", 2, "Specification",
     "Turn a design conversation into a published spec issue with /to-spec, and check the result before filing.",
     "Apply", "A2 (checker)", "mattpocock/skills: to-spec"),
    ("M10", 2, "Decomposition",
     "Decompose a spec into 6-10 tracer-bullet tickets with declared blocking edges using /to-tickets.",
     "Create", "A2 (rubric)", "mattpocock/skills: to-tickets"),
    ("M11", 2, "Work dependencies",
     "Record work dependencies as GitHub issue dependencies via the gh CLI, distinguish them from sub-issues and say why a sub-issue cannot express a two-parent join, and render the resulting graph as a Mermaid diagram.",
     "Apply", "A2 (checker)", "GitHub Docs: Creating issue dependencies; Adding sub-issues; Issue dependencies REST API"),
    ("M12", 2, "PM integration",
     "Connect Linear to the Instance and verify that issues and PRs appear and update without manual touching.",
     "Apply", "A2 (checker)", "Linear Docs: GitHub integration"),

    # --- Week 3 ---
    ("M13", 3, "Branching",
     "Create a branch bound to an issue with gh issue develop, and follow a branch-naming convention that auto-links in both GitHub and Linear.",
     "Apply", "A3 (checker)", "GitHub CLI manual: gh issue develop"),
    ("M14", 3, "Diff craft",
     "Judge whether a diff is reviewable - single purpose, bounded size, no drive-by changes - and split one that is not.",
     "Evaluate", "A3 (rubric); Quiz 1", "Course exercise"),
    ("M15", 3, "PR authoring",
     "Write a PR body that states intent, scope, verification, and risk, and closes its issue on merge with a linking keyword.",
     "Create", "A3 (checker + rubric)", "GitHub Docs: Linking a pull request to an issue"),
    ("M16", 3, "Stacking",
     "Build a stacked PR on a dependent branch, set its base correctly, and retarget it after its parent merges.",
     "Apply", "A3 (checker)", "GitHub Docs: Changing the base branch of a pull request"),
    ("M17", 3, "Conflicts",
     "Resolve a merge conflict on a stacked branch with /resolving-merge-conflicts and explain what caused it.",
     "Apply", "A3 (checker)", "mattpocock/skills: resolving-merge-conflicts"),
    ("M18", 3, "Toolbelt",
     "Choose the right skill for a situation using /ask-matt, and say what wayfinder, request-refactor-plan, grill-me, tdd, handoff, and git-guardrails each do.",
     "Understand", "Quiz 1; Self-Check", "mattpocock/skills repo"),

    # --- Week 4 ---
    ("M19", 4, "Review mechanics",
     "Distinguish a single comment, a review, and a suggested change; batch comments and submit a review with a verdict.",
     "Apply", "A4 (checker)", "GitHub Docs: Reviewing proposed changes in a pull request"),
    ("M20", 4, "Review mechanics",
     "Resolve conversations, re-request review, and read a PR's review state correctly.",
     "Apply", "A4 (checker)", "GitHub Docs: Commenting on a pull request"),
    ("M21", 4, "Agentic review",
     "Run /code-review on a branch and separate verified findings from noise across its Standards and Spec axes.",
     "Analyze", "A4 (written)", "code.claude.com: Code Review"),
    ("M22", 4, "Communication",
     "Write review comments a peer can act on without asking a clarifying question: located, specific, and severity-marked.",
     "Create", "A4 paired (rubric)", "Course exercise"),
    ("M23", 4, "Authoring",
     "Respond to review as the author: direct Claude to address findings, push, and re-request review without losing the thread.",
     "Apply", "A4 paired (checker)", "Course exercise"),
    ("M24", 4, "Bug capture",
     "Turn what review uncovers into filed issues with /qa rather than fixing out of scope in the open PR.",
     "Apply", "A4 (checker)", "mattpocock/skills: qa"),

    # --- Week 5 ---
    ("M25", 5, "CI",
     "Write a GitHub Actions workflow that runs pytest and ruff on every pull request, and diagnose a failing run from its logs.",
     "Apply", "A5 (checker)", "GitHub Docs: About continuous integration with GitHub Actions"),
    ("M26", 5, "Gates",
     "Configure a ruleset requiring status checks, review approval, and conversation resolution - and explain why a Free org must use a public repo for it to apply.",
     "Apply", "A5 (checker)", "GitHub Docs: About rulesets; About protected branches"),
    ("M27", 5, "Gates",
     "Route review with CODEOWNERS and org teams, and predict which reviewers a given PR will request.",
     "Apply", "A5 (checker)", "GitHub Docs: About code owners"),
    ("M28", 5, "Merging",
     "Choose merge vs squash vs rebase for a given situation and justify it in terms of history, revert, and bisect.",
     "Evaluate", "Quiz 2; A5 (rubric)", "GitHub Docs: About merge methods on GitHub"),
    ("M29", 5, "Merging",
     "Enable and use a merge queue, explaining what it protects against that required status checks alone do not.",
     "Understand", "Quiz 2; A5 (checker)", "GitHub Docs: Managing a merge queue"),
    ("M30", 5, "Agentic CI",
     "Read a working Claude Code GitHub Action workflow - triggers, permissions, secret, actor checks - and explain why it cannot run on a pull request from a fork.",
     "Analyze", "Quiz 2; A5 (written)", "code.claude.com: GitHub Actions; GitHub Docs: Using secrets in GitHub Actions"),

    # --- Week 6 ---
    ("M31", 6, "Forks",
     "Contribute to the Upstream Repo from a fork: sync the fork, open a cross-repo PR, and follow the repo's contribution etiquette.",
     "Apply", "A6 (checker)", "GitHub Docs: About forks; Contributing to a project"),
    ("M32", 6, "Maintainer seat",
     "Triage incoming external PRs and issues through the five-label state machine with /triage, writing agent-ready briefs.",
     "Apply", "A6 (rubric)", "mattpocock/skills: triage"),
    ("M33", 6, "PM integration",
     "Contrast Linear's model - teams, projects, cycles - with GitHub's repos, issues, and milestones, and name where the two disagree.",
     "Understand", "Quiz 2", "Linear Docs; GitHub Docs: About milestones"),
    ("M34", 6, "PM integration",
     "Drive issue status from git activity: branch names, magic words, and PR state transitions, with no manual status changes.",
     "Apply", "A6 (checker)", "Linear Docs: GitHub integration"),
    ("M35", 6, "Judgment",
     "Decide where the source of truth for work should live for a given team, and defend it against duplication, sync lag, and permissions.",
     "Evaluate", "Quiz 2; Capstone", "Course exercise"),
    ("M36", 6, "Synthesis",
     "Ship one complete loop end to end and write a Personal PR Workflow Playbook naming the gates and defaults you will carry to your own work.",
     "Create", "Capstone", "Course project"),
]

# Learning Pages per objective, as (path under docs/, title). The first entry is the
# objective's entry page (index.html); the rest are supporting pages. An objective with no
# entry has not been built yet. check_site.py validates the pages themselves; this list is
# what the workbook and its markdown mirror link to.
SITE_URL = "https://perro-ruidoso.github.io/gh-pr-mastery/"
PAGES = {
    "M01": [
        ("week-01/m01-instance-and-toolchain/index.html", "Your Instance and Your Toolchain"),
        ("week-01/m01-instance-and-toolchain/create-your-instance.html", "Create Your Instance"),
    ],
    "M02": [
        ("week-01/m02-what-a-pull-request-is/index.html", "Two Refs and a Conversation"),
        ("week-01/m02-what-a-pull-request-is/heads-forks-and-drafts.html", "Heads, Forks, and Drafts"),
    ],
    "M03": [
        ("week-01/m03-the-merge-base/index.html", "The Merge Base"),
        ("week-01/m03-the-merge-base/what-github-shows.html", "What GitHub Shows You"),
    ],
    "M04": [
        ("week-01/m04-reading-work-from-the-terminal/index.html", "Listing and Viewing"),
        ("week-01/m04-reading-work-from-the-terminal/filter-and-diff.html", "Filter and Diff"),
    ],
    "M05": [
        ("week-01/m05-configuring-the-skills/index.html", "Teaching Your Repo to the Skills"),
        ("week-01/m05-configuring-the-skills/labels-and-domain-docs.html", "Labels and Domain Docs"),
    ],
    "M06": [
        ("week-01/m06-reading-the-exhibit/index.html", "Reading the Exhibit"),
        ("week-01/m06-reading-the-exhibit/the-method.html", "The Method"),
        ("week-01/m06-reading-the-exhibit/your-target.html", "Your Target"),
    ],
    "M07": [
        ("week-02/m07-issues-an-agent-can-execute/index.html", "The Ticket as a Contract"),
        ("week-02/m07-issues-an-agent-can-execute/acceptance-criteria.html", "Criteria That Can Fail"),
    ],
    "M08": [
        ("week-02/m08-labels-types-and-templates/index.html", "Labels and Issue Types"),
        ("week-02/m08-labels-types-and-templates/issue-templates.html", "Issue Templates"),
    ],
    "M09": [
        ("week-02/m09-conversation-to-spec/index.html", "From Conversation to Spec"),
        ("week-02/m09-conversation-to-spec/checking-the-spec.html", "Checking the Spec"),
    ],
    "M10": [
        ("week-02/m10-decomposing-into-tickets/index.html", "Tracer Bullets and Blocking Edges"),
        ("week-02/m10-decomposing-into-tickets/the-seed-backlog.html", "Reading the Seed Backlog"),
    ],
    "M11": [
        ("week-02/m11-recording-dependencies/index.html", "Two Features, Two Shapes"),
        ("week-02/m11-recording-dependencies/reading-the-graph-back.html", "Reading the Graph Back"),
    ],
    "M12": [
        ("week-02/m12-wiring-linear/index.html", "Linear Above GitHub"),
        ("week-02/m12-wiring-linear/verify-without-touching.html", "Verify Without Touching"),
    ],
    "M13": [
        ("week-03/m13-branch-per-issue/index.html", "A Branch Bound to an Issue"),
        ("week-03/m13-branch-per-issue/naming-that-links.html", "Naming That Links Twice"),
    ],
    "M14": [
        ("week-03/m14-reviewable-diffs/index.html", "What a Reviewable Diff Looks Like"),
        ("week-03/m14-reviewable-diffs/splitting-a-drive-by.html", "Splitting a Drive-by"),
    ],
    "M15": [
        ("week-03/m15-the-pr-body/index.html", "Intent, Scope, Verification, Risk"),
        ("week-03/m15-the-pr-body/the-closing-keyword.html", "The Keyword and Its Two Conditions"),
    ],
    "M16": [
        ("week-03/m16-stacked-prs/index.html", "Stacking on a Branch That Is Not Main"),
        ("week-03/m16-stacked-prs/retargeting.html", "Retarget, Then Look Again"),
        ("week-03/m16-stacked-prs/deleting-the-parent.html", "What Deleting the Parent Does"),
    ],
    "M17": [
        ("week-03/m17-resolving-conflicts/index.html", "Where the Conflict Came From"),
        ("week-03/m17-resolving-conflicts/running-the-skill.html", "Running /resolving-merge-conflicts"),
    ],
    "M18": [
        ("week-03/m18-the-toolbelt/index.html", "Ask Matt First"),
        ("week-03/m18-the-toolbelt/six-skills-six-situations.html", "Six Skills, Six Situations"),
    ],
}

HEADERS = ["ID", "Week", "Unit", "Category", "Mastery Objective",
           "Bloom Level", "Assessed By", "Primary Source", "Learning Pages"]
WIDTHS = [7, 7, 34, 18, 82, 13, 22, 46, 60]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
WEEK_FILLS = ["EDF2FA", "FFFFFF"]
BLOOM_ORDER = ["Understand", "Apply", "Analyze", "Evaluate", "Create"]


def pages_cell(oid: str) -> str:
    """Multi-line cell text: one 'Title - docs/path' line per Learning Page, or a dash."""
    pages = PAGES.get(oid, [])
    if not pages:
        return "-"
    return "\n".join(f"{title} - docs/{path}" for path, title in pages)


def rows():
    for oid, week, category, objective, bloom, assessed, source in OBJECTIVES:
        yield [oid, week, UNITS[week], category, objective, bloom, assessed, source,
               pages_cell(oid)]


def build_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mastery Objectives"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows():
        ws.append(row)

    for row in ws.iter_rows(min_row=2):
        fill = PatternFill("solid", fgColor=WEEK_FILLS[row[1].value % 2])
        for cell in row:
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        pages = PAGES.get(row[0].value, [])
        if pages:
            link_cell = row[len(HEADERS) - 1]
            link_cell.hyperlink = SITE_URL + pages[0][0]
            link_cell.font = Font(color="0563C1", underline="single")

    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)


def bloom_counts() -> dict:
    counts = {}
    for obj in OBJECTIVES:
        bloom = obj[4]
        counts[bloom] = counts.get(bloom, 0) + 1
    return counts


def pages_md(oid: str) -> str:
    """Markdown links, relative to objectives/, so they resolve when read on GitHub. The
    first link is the objective's entry page; the others are its supporting pages."""
    pages = PAGES.get(oid, [])
    if not pages:
        return "not yet built"
    return "<br>".join(f"[{title}](../docs/{path})" for path, title in pages)


def build_md(path: Path) -> None:
    built = sum(1 for o in OBJECTIVES if PAGES.get(o[0]))
    page_count = sum(len(v) for v in PAGES.values())
    lines = [
        "# Mastery Objectives - GitHub Pull Requests for Agentic Coders",
        "",
        "Generated by `objectives/build_objectives.py`. Edit the Python, not this file.",
        "",
        f"**Learning Pages:** {built} of {len(OBJECTIVES)} objectives have pages ({page_count} pages). "
        "Each objective's first link is its entry page (`index.html`); the rest are supporting "
        f"pages. The site is served at <{SITE_URL}>; the links below are repo-relative.",
        "",
    ]
    for week, unit in UNITS.items():
        lines += [
            f"## Week {week} - {unit}",
            "",
            "| ID | Category | Objective | Bloom | Assessed By | Primary Source | Learning Pages |",
            "|---|---|---|---|---|---|---|",
        ]
        for oid, w, category, objective, bloom, assessed, source in OBJECTIVES:
            if w == week:
                lines.append(
                    f"| {oid} | {category} | {objective} | {bloom} | {assessed} | {source} "
                    f"| {pages_md(oid)} |"
                )
        lines.append("")

    counts = bloom_counts()
    lines += ["## Bloom distribution", "", "| Level | Count |", "|---|---|"]
    for level in BLOOM_ORDER:
        lines.append(f"| {level} | {counts.get(level, 0)} |")
    lines.append("")

    path.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    assert len(OBJECTIVES) == 36, f"expected 36 objectives, found {len(OBJECTIVES)}"
    ids = [o[0] for o in OBJECTIVES]
    expected = [f"M{n:02d}" for n in range(1, 37)]
    assert ids == expected, "objective IDs are not M01..M36 in order"
    unknown = {o[4] for o in OBJECTIVES} - set(BLOOM_ORDER)
    assert not unknown, f"unknown Bloom levels: {unknown}"
    assert set(PAGES) <= set(ids), f"PAGES names unknown objectives: {set(PAGES) - set(ids)}"
    docs = HERE.parent / "docs"
    for oid, pages in PAGES.items():
        assert pages and pages[0][0].endswith("/index.html"), f"{oid}: first page must be index.html"
        for rel, _title in pages:
            assert (docs / rel).exists(), f"{oid}: listed page does not exist: docs/{rel}"
    build_xlsx(HERE / "mastery-objectives.xlsx")
    build_md(HERE / "mastery-objectives.md")
    print(f"wrote {len(OBJECTIVES)} objectives to mastery-objectives.xlsx and .md")
    print(f"learning pages: {sum(len(v) for v in PAGES.values())} across {len(PAGES)} objectives")
    print("bloom:", bloom_counts())
